"""Generación del Excel de salida: datos de entrada + resultados + gráfico embebido."""

from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference

from ..schemas import CalculoResponse

_NOMBRES_MODELOS = {
    "ley_de_masas": "Ley de masas",
    "ley_de_masas_corregida": "Ley de masas corregida",
    "sharp": "Sharp",
    "iso12354": "ISO 12354-1",
    "davy": "Davy",
}


def generar_excel(datos: CalculoResponse) -> io.BytesIO:
    wb = Workbook()

    # --- Hoja de datos de entrada ---
    ws_entrada = wb.active
    ws_entrada.title = "Datos de entrada"
    ws_entrada.append(["Parámetro", "Valor"])
    for campo, valor in datos.entrada.model_dump().items():
        ws_entrada.append([campo, valor if valor is not None else ""])

    # --- Hoja de resultados ---
    ws_res = wb.create_sheet("Resultados")
    modelos_disponibles = [
        clave for clave, valores in datos.resultados.model_dump().items() if valores is not None
    ]
    encabezado = ["Frecuencia (Hz)"] + [_NOMBRES_MODELOS[c] for c in modelos_disponibles]
    ws_res.append(encabezado)

    resultados_dict = datos.resultados.model_dump()
    for i, f in enumerate(datos.frecuencias):
        fila = [f] + [round(resultados_dict[c][i], 2) for c in modelos_disponibles]
        ws_res.append(fila)

    # --- Gráfico embebido ---
    if modelos_disponibles:
        chart = LineChart()
        chart.title = "R vs frecuencia"
        chart.x_axis.title = "Frecuencia (Hz)"
        chart.y_axis.title = "R (dB)"
        chart.style = 2

        # openpyxl no marca los ejes como visibles por default: sin esto,
        # Excel los abre sin las marcas de escala (números) en X e Y.
        chart.x_axis.delete = False
        chart.y_axis.delete = False
        chart.x_axis.tickLblPos = "nextTo"
        chart.y_axis.tickLblPos = "nextTo"
        # Además, openpyxl deja axPos="l" (izquierda) en los DOS ejes por
        # default. Con el eje de categorías (frecuencia) también en "l",
        # Excel no calcula bien el layout y termina sin dibujar los
        # números de ninguno de los dos ejes. El de categorías va abajo.
        chart.x_axis.axPos = "b"
        chart.y_axis.axPos = "l"

        n_filas = len(datos.frecuencias)
        cats = Reference(ws_res, min_col=1, min_row=2, max_row=n_filas + 1)
        for idx, clave in enumerate(modelos_disponibles):
            col = idx + 2
            data = Reference(ws_res, min_col=col, min_row=1, max_row=n_filas + 1)
            chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        ws_res.add_chart(chart, f"{chr(ord('A') + len(encabezado) + 1)}2")

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
